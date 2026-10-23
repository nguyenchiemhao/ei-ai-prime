"""Generators of small test documents (PDF, DOCX) shared by the ingestion tests."""

import zipfile
from io import BytesIO

from docx import Document
from pypdf import PdfReader, PdfWriter

ZIP_BOMB_BYTES = 260 * 1024 * 1024
_ZIP_CHUNK = b"\0" * (1024 * 1024)
_APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    "<Pages>{pages}</Pages></Properties>"
)


def _pdf_escape(line: str) -> str:
    """Escapes the characters that are special inside a PDF string."""
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_text_pdf(pages: list[str]) -> bytes:
    """A PDF with a Helvetica text layer: one page per entry, ``\\n`` starts a new line.

    The text must be ASCII.
    """
    objects: list[bytes] = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids: list[str] = []
    for text in pages:
        lines = [_pdf_escape(line) for line in text.split("\n")]
        stream = "BT /F1 10 Tf 12 TL 20 780 Td " + " T* ".join(f"({ln}) Tj" for ln in lines) + " ET"
        content = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode("ascii")
        page_number = len(objects) + 1
        kids.append(f"{page_number} 0 R")
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_number + 1} 0 R >>"
            ).encode("ascii")
        )
        objects.append(content)
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode("ascii")
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii")
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    size = len(objects) + 1
    trailer = f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    out += trailer.encode("ascii")
    return bytes(out)


def make_blank_pdf(pages: int) -> bytes:
    """A PDF of ``pages`` empty pages (what a scan without a text layer looks like to a parser)."""
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_encrypted_pdf() -> bytes:
    """A one-page PDF that needs the password ``secret`` (RC4, so no extra crypto library)."""
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.encrypt("secret", algorithm="RC4-128")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_aes_pdf(user_password: str, text: str = "Secret page text " * 5) -> bytes:
    """A one-page text PDF encrypted with AES-256; ``user_password`` may be empty."""
    writer = PdfWriter(clone_from=PdfReader(BytesIO(make_text_pdf([text]))))
    writer.encrypt(user_password, owner_password="owner", algorithm="AES-256")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(
    blocks: list[tuple[str, str]],
    table: list[list[str]] | None = None,
    pages: int | None = None,
) -> bytes:
    """A DOCX with one paragraph per ``(style, text)`` entry, then the optional table.

    ``pages`` is written to ``docProps/app.xml`` as ``<Pages>``.
    """
    document = Document()
    for style, text in blocks:
        document.add_paragraph(text, style=style)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = BytesIO()
    document.save(buffer)
    if pages is None:
        return buffer.getvalue()
    source = zipfile.ZipFile(BytesIO(buffer.getvalue()))
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename != "docProps/app.xml":
                target.writestr(info.filename, source.read(info.filename))
        target.writestr("docProps/app.xml", _APP_XML.format(pages=pages))
    return result.getvalue()


def add_docx_part(data: bytes, name: str, size: int, override: str | None = None) -> bytes:
    """The DOCX with an extra part ``name`` of ``size`` bytes (spaces, deflated).

    ``override`` registers a content type for that part in ``[Content_Types].xml``.
    """
    source = zipfile.ZipFile(BytesIO(data))
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if override and info.filename == "[Content_Types].xml":
                line = f'<Override PartName="/{name}" ContentType="{override}"/></Types>'
                content = content.replace(b"</Types>", line.encode("ascii"))
            target.writestr(info.filename, content)
        with target.open(name, "w", force_zip64=True) as entry:
            left = size
            while left > 0:
                piece = min(left, len(_ZIP_CHUNK))
                entry.write(b" " * piece)
                left -= piece
    return result.getvalue()


def make_zip_bomb_docx() -> bytes:
    """A small DOCX-named ZIP whose entries expand to more than 250 MB."""
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        with archive.open("word/bomb.bin", "w", force_zip64=True) as entry:
            for _ in range(ZIP_BOMB_BYTES // len(_ZIP_CHUNK)):
                entry.write(_ZIP_CHUNK)
    return result.getvalue()
